import os
import re
import uuid
from datetime import datetime, timezone

from docx import Document
from docx.text.paragraph import Paragraph


def _is_section_header(paragraph: Paragraph) -> bool:
    text = paragraph.text.strip()
    if not text:
        return False
    style_name = paragraph.style.name if paragraph.style else ""
    if style_name in ("Heading 1", "Heading 2", "Heading 3"):
        return True
    if text.isupper() and len(text) < 80:
        return True
    if len(text) < 80 and not text.endswith("."):
        for run in paragraph.runs:
            if run.bold and run.text.strip():
                return True
    return False


def extract_docx_sections(file_path: str) -> list[dict]:
    doc = Document(file_path)
    sections: list[dict] = []
    current_title = ""
    current_content: list[str] = []

    def flush_section() -> None:
        if current_title or current_content:
            content = "\n".join(current_content).strip()
            if content or current_title:
                sections.append(
                    {"section_title": current_title or "Untitled", "content": content}
                )

    for element in doc.element.body:
        tag = element.tag.split("}")[-1] if "}" in element.tag else element.tag
        if tag == "p":
            para = Paragraph(element, doc)
            text = para.text.strip()
            if not text:
                continue
            if _is_section_header(para):
                flush_section()
                current_title = text
                current_content = []
            else:
                current_content.append(text)
        elif tag == "tbl":
            from docx.table import Table

            table = Table(element, doc)
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    current_content.append(" | ".join(cells))

    flush_section()

    if not sections:
        all_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                all_text.append(para.text.strip())
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    all_text.append(" | ".join(cells))
        content = "\n".join(all_text).strip()
        if content:
            sections.append({"section_title": "Full document", "content": content})

    return sections


def chunk_section(
    section_title: str,
    content: str,
    max_chars: int = 800,
    overlap_chars: int = 100,
) -> list[str]:
    if not content.strip():
        return []

    sentences = re.split(r"(?<=[.!?])\s+|\.\n", content)
    sentences = [s.strip() for s in sentences if s.strip()]

    chunks: list[str] = []
    current = ""

    for sentence in sentences:
        candidate = f"{current} {sentence}".strip() if current else sentence
        if len(candidate) <= max_chars:
            current = candidate
        else:
            if current:
                chunk_text = f"[{section_title}]\n{current}"
                if len(current) >= 50:
                    chunks.append(chunk_text)
                overlap = current[-overlap_chars:] if len(current) > overlap_chars else current
                current = f"{overlap} {sentence}".strip()
            else:
                current = sentence[:max_chars]

    if current:
        chunk_text = f"[{section_title}]\n{current}"
        if len(current) >= 50:
            chunks.append(chunk_text)

    return chunks


def ingest_policy_document(
    file_path: str,
    doc_title: str,
    description: str,
    db_repo,
    vector_store,
) -> dict:
    doc_id = f"policy_{uuid.uuid4().hex[:8]}"
    filename = os.path.basename(file_path)
    sections = extract_docx_sections(file_path)

    chunk_dicts: list[dict] = []
    chunk_index = 0
    for section in sections:
        section_chunks = chunk_section(section["section_title"], section["content"])
        for chunk_text in section_chunks:
            chunk_dicts.append(
                {
                    "text": chunk_text,
                    "chunk_index": chunk_index,
                    "section_title": section["section_title"],
                    "doc_title": doc_title,
                    "filename": filename,
                }
            )
            chunk_index += 1

    total_chunks = vector_store.ingest_policy_chunks(doc_id, chunk_dicts)

    db_repo.upsert_policy_doc(
        {
            "doc_id": doc_id,
            "filename": filename,
            "title": doc_title,
            "description": description,
            "chunk_count": total_chunks,
            "status": "active",
            "uploaded_at": datetime.now(timezone.utc).isoformat(),
        }
    )

    return {
        "doc_id": doc_id,
        "title": doc_title,
        "chunk_count": total_chunks,
        "sections_found": len(sections),
    }


def ingest_policies_from_directory(
    policies_dir: str, db_repo, vector_store
) -> int:
    if not os.path.isdir(policies_dir):
        return 0

    count = 0
    for filename in os.listdir(policies_dir):
        if not filename.lower().endswith(".docx"):
            continue
        existing = db_repo.get_policy_doc_by_filename(filename)
        if existing:
            continue
        file_path = os.path.join(policies_dir, filename)
        title = os.path.splitext(filename)[0]
        ingest_policy_document(file_path, title, "", db_repo, vector_store)
        count += 1
    return count
